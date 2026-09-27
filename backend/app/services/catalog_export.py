"""The catalogue, once, in the shape each marketplace asks for.

Listing on Amazon, Flipkart, Meesho, Myntra or AJIO means filling that
portal's bulk-upload template: a category-specific Excel file downloaded
from the seller panel, with its own column names, changed without notice.
Typing the same twelve pieces into five of them is where catalogues drift
apart - a price corrected on one and not the others, a size missing on one.

So there is one source - the pieces as ZISUN records them - and two outputs:

* **The master sheet** (`master_xlsx`): one row per size and colour with
  every field ZISUN holds, image links included. Always correct, because
  it is only our own data.
* **A filled template** (`inspect_template` then `fill_template`): she
  uploads the marketplace's blank template; the header row is found by
  matching its column names against the aliases below; the matches are
  shown for her to check or override; and the same file comes back with
  one row per variant written under the header. Reading the real template
  every time is what keeps this working when a marketplace renames a
  column - a guessed template would be rejected on upload.

The aliases are the spellings seen in the portals' templates and are
matched loosely (case, spaces and punctuation ignored). Anything not
matched is left blank for her, never guessed: a wrong value in a
marketplace listing is worse than an empty one.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from typing import Any, Iterable, Optional

BRAND = "ZISUN"

# key, label for the master sheet, template header aliases (normalised)
FIELDS: list[tuple[str, str, list[str]]] = [
    ("sku", "SKU", ["sku", "sellersku", "itemsku", "sellerskuid", "skuid", "skucode", "sellerskucode",
                    "vendorskucode", "vendorsku", "styleskucode", "merchantsku", "productsku"]),
    ("parent_sku", "Parent SKU (one per piece)", ["parentsku", "parentskuid", "stylegroupid", "stylecode",
                                                  "parentstylecode", "groupid", "catalogname"]),
    ("title", "Title", ["itemname", "title", "producttitle", "productname", "stylename", "name", "productdisplayname"]),
    ("variant_title", "Title with size and colour", ["variantname", "variantitle", "itemnamewithvariant"]),
    ("brand", "Brand", ["brand", "brandname"]),
    ("description", "Description", ["productdescription", "description", "longdescription", "styledescription"]),
    ("bullet_1", "Key feature 1", ["bulletpoint1", "keyfeature1", "keyfeatures1", "highlight1", "feature1"]),
    ("bullet_2", "Key feature 2", ["bulletpoint2", "keyfeature2", "keyfeatures2", "highlight2", "feature2"]),
    ("bullet_3", "Key feature 3", ["bulletpoint3", "keyfeature3", "keyfeatures3", "highlight3", "feature3"]),
    ("bullet_4", "Key feature 4", ["bulletpoint4", "keyfeature4", "keyfeatures4", "highlight4", "feature4"]),
    ("bullet_5", "Key feature 5", ["bulletpoint5", "keyfeature5", "keyfeatures5", "highlight5", "feature5"]),
    ("gender", "Gender", ["gender", "targetgender", "departmentname", "department", "idealfor"]),
    ("age_group", "Age group", ["agegroup", "targetagegroup", "agerangedescription"]),
    ("category", "ZISUN category", ["category", "subcategory", "producttype", "articletype", "itemtypename"]),
    ("colour", "Colour", ["colour", "color", "colorname", "colourname", "primarycolor", "primarycolour",
                          "basecolour", "basecolor", "colorfamily"]),
    ("size", "Size", ["size", "sizename", "brandsize", "standardsize", "apparelsize", "sizevalue"]),
    ("mrp", "MRP (₹, incl. GST)", ["mrp", "maximumretailprice", "listprice", "mrpinr", "productmrp"]),
    ("price", "Selling price (₹, incl. GST)", ["standardprice", "sellingprice", "yoursellingprice", "price",
                                               "sellingpriceinr", "offerprice", "yourprice", "saleprice"]),
    ("gst_rate", "GST rate (%)", ["gst", "gstrate", "gstpercent", "taxrate", "gstpercentage", "producttaxcode"]),
    ("hsn", "HSN code", ["hsn", "hsncode", "hsnsac"]),
    ("quantity", "Stock", ["quantity", "stock", "inventory", "availablequantity", "sellablequantity", "qty",
                           "stockcount", "inventorycount"]),
    ("fabric", "Fabric", ["fabric", "fabrictype", "material", "materialtype", "fabriccomposition", "outermaterial",
                          "topfabric", "materialcomposition"]),
    ("pattern", "Pattern", ["pattern", "patterntype", "patternname", "print", "printorpatterntype", "printtype"]),
    ("sleeve", "Sleeve", ["sleevelength", "sleevetype", "sleeve", "sleevestyling"]),
    ("neck", "Neck", ["neck", "necktype", "neckstyle", "neckline"]),
    ("fit", "Fit", ["fit", "fittype", "fitsilhouette"]),
    ("length", "Length", ["length", "garmentlength", "topslength", "itemlength", "kurtalength"]),
    ("occasion", "Occasion", ["occasion", "occasiontype", "usage", "specialfeatureoccasion"]),
    ("embroidery", "Embroidery / ornamentation", ["embroidery", "ornamentation", "embellishment", "surfacestyling"]),
    ("bottom_type", "Bottom type", ["bottomtype", "bottomwear", "bottomstyle"]),
    ("dupatta", "Dupatta included", ["dupatta", "dupattaincluded", "withdupatta"]),
    ("set_contents", "In the set", ["setcontents", "packcontains", "setcontains", "includedcomponents"]),
    # A count, not the list: "Pack of" and "Number of items" columns take a
    # number, and "Kurta, Palazzo" in one is a rejected row.
    ("pieces_count", "Pieces in the set", ["packof", "numberofitems", "numberofpieces", "unitcount", "itemcount"]),
    ("wash_care", "Wash care", ["washcare", "careinstructions", "caretype", "washinstructions"]),
    ("net_quantity", "Net quantity", ["netquantity", "itempackagequantity", "netqty"]),
    ("country_of_origin", "Country of origin", ["countryoforigin", "countryoforiginname", "origincountry"]),
    ("manufacturer", "Manufacturer", ["manufacturer", "manufacturername", "manufacturerdetails", "manufacturercontactinformation"]),
    ("manufacturer_address", "Manufacturer address", ["manufactureraddress"]),
    ("packer", "Packer", ["packer", "packerdetails", "packername", "packercontactinformation"]),
    ("commodity", "Generic name", ["genericname", "commodityname", "itemtype", "commodity"]),
] + [
    (f"image_{i}", f"Image {i}", ([f"mainimageurl", "mainimage", "image1", "imageurl1", "frontimage", "primaryimage",
                                    "mainimageurl1", "image1url", "front"] if i == 1 else
                                   [f"otherimageurl{i - 1}", f"image{i}", f"imageurl{i}", f"image{i}url",
                                    f"additionalimage{i - 1}", f"otherimage{i - 1}"]))
    for i in range(1, 9)
]

REQUIRED_FOR_TEMPLATE = ("sku", "title", "price")
FIELD_LABEL = {k: label for k, label, _ in FIELDS}


def _norm(s: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


def _yes_no(v: Optional[bool]) -> str:
    return "" if v is None else ("Yes" if v else "No")


def _rupees(paise: Optional[int]) -> Optional[float]:
    return None if paise is None else round(paise / 100, 2)


# ── Our data -> rows ─────────────────────────────────────────────────────────

def _bullets(p) -> list[str]:
    out = []
    fabric = " ".join(x for x in [p.fabric_composition, (p.weave or "")] if x).strip()
    if fabric:
        out.append(f"Fabric: {fabric}" + (f", {p.fabric_gsm} GSM" if p.fabric_gsm else ""))
    shape = ", ".join(x for x in [p.fit, p.garment_length] if x)
    if shape:
        out.append(f"Fit and length: {shape}")
    detail = ", ".join(x for x in [p.neck_type, p.sleeve_type] if x)
    if detail:
        out.append(f"Neck and sleeve: {detail}")
    pieces = [x for x in (p.set_pieces or []) if x]
    if pieces:
        out.append("In the set: " + ", ".join(pieces))
    if p.wash_care:
        out.append(f"Care: {p.wash_care}")
    return out[:5]


def _images(p, variant_id) -> list[str]:
    media = [m for m in (p.media or []) if str(getattr(m.type, "value", m.type)).upper() == "IMAGE"]
    media.sort(key=lambda m: (m.display_order or 0))
    tagged = [m for m in media if m.variant_id == variant_id]
    untagged = [m for m in media if m.variant_id is None]
    seen, out = set(), []
    for m in tagged + untagged:
        url = m.cdn_url or m.url
        if url and url not in seen:
            seen.add(url)
            out.append(url)
    return out[:8]


def rows_for(products: Iterable, *, legal, rate_for, default_hsn: str) -> list[dict]:
    """One row per live size/colour. `legal(product)` returns the resolved
    Legal Metrology block and `rate_for(paise)` the GST rate - passed in so
    this stays the same arithmetic the storefront and invoices use."""
    rows: list[dict] = []
    for p in products:
        lm = legal(p)
        parent = _norm(p.name)[:18].upper() or str(p.id)[:8]
        bullets = _bullets(p)
        for v in sorted(p.variants or [], key=lambda v: (v.color or "", v.size or "")):
            if not getattr(v, "is_active", True):
                continue
            price = (p.base_price or 0) + (v.price_delta or 0)
            mrp = p.compare_at_price if (p.compare_at_price and p.compare_at_price > price) else price
            colour = v.color or p.colour
            row = {
                "sku": v.sku, "parent_sku": f"{BRAND}-{parent}",
                "title": p.name,
                "variant_title": f"{p.name} - {', '.join(x for x in [colour, v.size] if x)}",
                "brand": BRAND, "description": (p.description or "").strip(),
                "gender": "Women", "age_group": "Adult",
                "category": p.category.name if getattr(p, "category", None) else "",
                "colour": colour or "", "size": v.size or "",
                "mrp": _rupees(mrp), "price": _rupees(price),
                "gst_rate": rate_for(price), "hsn": p.hsn_code or default_hsn,
                "quantity": v.stock,
                "fabric": p.fabric_composition or "", "pattern": p.pattern or p.print_type or "",
                "sleeve": p.sleeve_type or "", "neck": p.neck_type or "", "fit": p.fit or "",
                "length": p.garment_length or "", "occasion": p.occasion or "",
                "embroidery": p.embroidery or "", "bottom_type": p.bottom_type or "",
                "dupatta": _yes_no(p.dupatta_included),
                "set_contents": ", ".join(x for x in (p.set_pieces or []) if x),
                "pieces_count": len([x for x in (p.set_pieces or []) if x]) or 1,
                "wash_care": p.wash_care or "",
                "net_quantity": lm.net_quantity, "country_of_origin": lm.country_of_origin,
                "manufacturer": lm.manufacturer_name, "manufacturer_address": lm.manufacturer_address,
                "packer": f"{lm.manufacturer_name}, {lm.manufacturer_address}",
                "commodity": lm.commodity_name,
            }
            for i in range(5):
                row[f"bullet_{i + 1}"] = bullets[i] if i < len(bullets) else ""
            imgs = _images(p, v.id)
            for i in range(8):
                row[f"image_{i + 1}"] = imgs[i] if i < len(imgs) else ""
            rows.append(row)
    return rows


# ── The master sheet ─────────────────────────────────────────────────────────

def master_xlsx(rows: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "ZISUN catalogue"
    keys = [k for k, _, _ in FIELDS]
    ws.append([FIELD_LABEL[k] for k in keys])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="F2F2F2")
    for r in rows:
        ws.append([r.get(k, "") for k in keys])
    ws.freeze_panes = "B2"
    for i, k in enumerate(keys, start=1):
        width = 60 if k in ("description",) or k.startswith("image_") else 24 if k.startswith("bullet") or k == "title" else 14
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def master_csv(rows: list[dict]) -> bytes:
    keys = [k for k, _, _ in FIELDS]
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([FIELD_LABEL[k] for k in keys])
    for r in rows:
        w.writerow([r.get(k, "") for k in keys])
    return ("﻿" + buf.getvalue()).encode("utf-8")   # BOM: Excel opens ₹ correctly


# ── Filling a marketplace's template ─────────────────────────────────────────

def match_header(header: Any) -> Optional[str]:
    """ZISUN field for a template column name, or None. Exact alias first;
    then an alias of 6+ characters contained in the header ("Product Title
    (max 100 chars)"). Short ones like "size" match exactly only, or
    "size chart" would take the size."""
    h = _norm(header)
    if not h:
        return None
    for key, _, aliases in FIELDS:
        if h in aliases:
            return key
    best = None
    for key, _, aliases in FIELDS:
        for a in aliases:
            if len(a) >= 6 and a in h and (best is None or len(a) > best[1]):
                best = (key, len(a))
    return best[0] if best else None


@dataclass
class TemplateLayout:
    sheet: str
    header_row: int                              # 1-based
    columns: list[tuple[int, str, Optional[str]]] = field(default_factory=list)   # (col, header, field)
    first_data_row: int = 0

    @property
    def mapped(self) -> dict[int, str]:
        return {c: f for c, _, f in self.columns if f}


def _load(data: bytes, filename: str):
    from openpyxl import load_workbook
    keep_vba = filename.lower().endswith(".xlsm")
    return load_workbook(io.BytesIO(data), keep_vba=keep_vba)


def inspect_template(data: bytes, filename: str, overrides: Optional[dict[str, str]] = None) -> TemplateLayout:
    """Find the sheet and row that hold the column names: the (sheet, row) in
    the first 15 rows of each sheet whose cells match the most ZISUN fields.
    Instructions and example rows are skipped by the same rule."""
    wb = _load(data, filename)
    best: Optional[tuple[int, str, int]] = None
    for ws in wb.worksheets:
        for r in range(1, min(ws.max_row, 15) + 1):
            hits = sum(1 for c in ws[r] if match_header(c.value))
            if hits and (best is None or hits > best[0]):
                best = (hits, ws.title, r)
    if best is None:
        raise ValueError("No column in this file looks like a product field. Is it the marketplace's blank upload template?")
    _, sheet, hr = best
    ws = wb[sheet]
    over = {int(k): v for k, v in (overrides or {}).items() if str(k).isdigit()}
    cols = []
    used: set[str] = set()
    for c in ws[hr]:
        if c.value is None or str(c.value).strip() == "":
            continue
        f = over.get(c.column, match_header(c.value))
        if f == "":
            f = None
        # one ZISUN field may fill several columns (MRP twice is common); but
        # an automatic match never assigns the same field to two neighbours
        # with different meaning - overrides are hers to decide.
        if f and c.column not in over and f in used and not f.startswith(("mrp", "price")):
            f = None
        if f:
            used.add(f)
        cols.append((c.column, str(c.value).strip(), f))
    layout = TemplateLayout(sheet=sheet, header_row=hr, columns=cols)
    layout.first_data_row = _first_empty_row(ws, hr, [c for c, _, f in cols if f])
    return layout


def _first_empty_row(ws, header_row: int, mapped_cols: list[int]) -> int:
    """The first row after the header where every mapped column is empty -
    templates often carry a description or example row under the header."""
    r = header_row + 1
    while r <= ws.max_row + 1:
        if all(ws.cell(r, c).value in (None, "") for c in mapped_cols):
            return r
        r += 1
    return r


def fill_template(data: bytes, filename: str, rows: list[dict],
                  overrides: Optional[dict[str, str]] = None) -> tuple[bytes, TemplateLayout]:
    layout = inspect_template(data, filename, overrides)
    wb = _load(data, filename)
    ws = wb[layout.sheet]
    r = layout.first_data_row
    for row in rows:
        for col, f in layout.mapped.items():
            v = row.get(f)
            if v not in (None, ""):
                ws.cell(r, col).value = v
        r += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), layout
