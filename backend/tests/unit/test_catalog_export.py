"""Catalogue export: our rows, and marketplace templates filled correctly.

The templates here are built to mimic how the portals lay theirs out - an
instructions sheet, a settings row, two header rows, a description row
under the header, asterisked names, a dropdown - because those are the
things that make a naive "write from row 2" filler put data in the wrong
place.
"""
import io
from types import SimpleNamespace as NS

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from app.schemas.catalog import LegalMetrology
from app.services import catalog_export as ce
from app.services import gst


def _product():
    img = lambda url, order, vid=None: NS(url=url, cdn_url=None, type=NS(value="IMAGE"), display_order=order, variant_id=vid)
    v_m = NS(id="vm", sku="ZS-WIN-M", size="M", color="Wine", stock=3, price_delta=0, is_active=True)
    v_l = NS(id="vl", sku="ZS-WIN-L", size="L", color="Wine", stock=0, price_delta=5000, is_active=True)
    v_off = NS(id="vo", sku="ZS-WIN-S", size="S", color="Wine", stock=9, price_delta=0, is_active=False)
    return NS(
        id="p1", name="Rich Wine Dabu Kurta Set", description="Hand-block dabu print.", base_price=112400,
        compare_at_price=149900, category=NS(name="Kurta sets"), colour="Wine", hsn_code=None,
        fabric_composition="Cotton", weave=None, fabric_gsm=None, fit="Straight cut", garment_length="Knee length",
        neck_type="Round neck", sleeve_type="3/4th sleeves", set_pieces=["Kurta", "Palazzo", "Dupatta"],
        wash_care="Hand wash cold", pattern="Geometric", print_type="Bandhini", occasion="Festive",
        embroidery=None, bottom_type="Palazzo", dupatta_included=True,
        commodity_name=None, net_quantity=None, dimensions=None, country_of_origin=None, manufacturer_name=None,
        variants=[v_l, v_m, v_off],
        media=[img("https://cdn/x/2.jpg", 2), img("https://cdn/x/1.jpg", 1), img("https://cdn/x/L.jpg", 0, "vl"),
               NS(url="https://cdn/x/v.mp4", cdn_url=None, type=NS(value="VIDEO"), display_order=0, variant_id=None)],
    )


def _rows():
    return ce.rows_for([_product()], legal=LegalMetrology.resolve, rate_for=gst.rate_for, default_hsn=gst.DEFAULT_HSN)


def test_one_row_per_live_size_with_our_own_numbers():
    rows = _rows()
    assert [r["sku"] for r in rows] == ["ZS-WIN-L", "ZS-WIN-M"], "off-sale S is not exported"
    m = next(r for r in rows if r["size"] == "M")
    assert (m["price"], m["mrp"]) == (1124.0, 1499.0), "struck-through price becomes the MRP"
    l = next(r for r in rows if r["size"] == "L")
    assert (l["price"], l["mrp"]) == (1174.0, 1499.0)
    assert m["gst_rate"] == gst.rate_for(112400) and m["hsn"] == gst.DEFAULT_HSN
    assert m["pieces_count"] == 3 and m["set_contents"] == "Kurta, Palazzo, Dupatta"
    assert m["net_quantity"] == "1 set - 3 pieces"
    assert m["dupatta"] == "Yes" and m["brand"] == "ZISUN" and m["gender"] == "Women"
    assert m["bullet_1"].startswith("Fabric: Cotton") and m["bullet_4"] == "In the set: Kurta, Palazzo, Dupatta"


def test_images_put_the_sizes_own_photo_first_and_skip_video():
    rows = {r["size"]: r for r in _rows()}
    assert rows["L"]["image_1"] == "https://cdn/x/L.jpg"
    assert rows["L"]["image_2"] == "https://cdn/x/1.jpg" and rows["L"]["image_3"] == "https://cdn/x/2.jpg"
    assert rows["M"]["image_1"] == "https://cdn/x/1.jpg", "an untagged photo leads when the size has none"
    assert not any(r[f"image_{i}"].endswith(".mp4") for r in rows.values() for i in range(1, 9))


def _save(wb) -> bytes:
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def amazon_like() -> bytes:
    wb = Workbook()
    ins = wb.active; ins.title = "Instructions"
    ins["A1"] = "How to fill this template"; ins["A2"] = "Enter one product per row"
    t = wb.create_sheet("Template")
    t.append(["TemplateType=fptcustom", "Version=2026.0910", "Category=clothing"])
    t.append(["Seller SKU", "Product Name", "Brand Name", "Standard Price", "Maximum Retail Price",
              "Main Image URL", "Other Image URL1", "Size", "Colour", "Quantity", "Parentage"])
    t.append(["item_sku", "item_name", "brand_name", "standard_price", "maximum_retail_price",
              "main_image_url", "other_image_url1", "size_name", "color_name", "quantity", "parent_child"])
    return _save(wb)


def flipkart_like() -> bytes:
    wb = Workbook()
    t = wb.active; t.title = "kurta"
    t.append(["Seller SKU ID*", "Product Title (max 100 chars)", "MRP (INR)*", "Your selling price (INR)*",
              "Size*", "Size Chart", "Pack of", "Fabric", "HSN*", "Occasion", "Style Code"])
    t.append(["Unique id", "Name shown to buyers", "Incl. tax", "Incl. tax", "Choose", "Upload image",
              "Number", "Material", "8 digits", "Choose", "Your code"])
    dv = DataValidation(type="list", formula1='"XS,S,M,L,XL"'); t.add_data_validation(dv); dv.add("E3:E500")
    return _save(wb)


def test_amazon_style_header_rows_are_found_past_settings_and_instructions():
    layout = ce.inspect_template(amazon_like(), "clothing.xlsx")
    assert layout.sheet == "Template" and layout.header_row in (2, 3)
    mapped = {h: f for _, h, f in layout.columns}
    got = {f for f in mapped.values() if f}
    assert {"sku", "title", "brand", "price", "mrp", "image_1", "image_2", "size", "colour", "quantity"} <= got
    assert layout.first_data_row == 4, "never over the second header row"


def test_flipkart_style_description_row_is_skipped_and_size_chart_left_alone():
    layout = ce.inspect_template(flipkart_like(), "kurta.xlsx")
    by_header = {h: f for _, h, f in layout.columns}
    assert by_header["Seller SKU ID*"] == "sku"
    assert by_header["Product Title (max 100 chars)"] == "title"
    assert by_header["MRP (INR)*"] == "mrp" and by_header["Your selling price (INR)*"] == "price"
    assert by_header["Size*"] == "size" and by_header["Size Chart"] is None
    assert by_header["Pack of"] == "pieces_count", "a count, never 'Kurta, Palazzo'"
    assert by_header["HSN*"] == "hsn"
    assert layout.first_data_row == 3


def test_fill_writes_under_the_header_and_keeps_the_dropdown():
    out, layout = ce.fill_template(flipkart_like(), "kurta.xlsx", _rows())
    wb = load_workbook(io.BytesIO(out))
    t = wb["kurta"]
    assert t["A2"].value == "Unique id", "the description row is untouched"
    assert [t["A3"].value, t["A4"].value] == ["ZS-WIN-L", "ZS-WIN-M"]
    assert t["C4"].value == 1499.0 and t["D4"].value == 1124.0
    assert t["E4"].value == "M" and t["F4"].value is None and t["G4"].value == 3
    assert t.data_validations.dataValidation, "the size dropdown survived"


def test_her_choice_overrides_the_match_and_blank_means_leave_empty():
    col_of = {h: c for c, h, _ in ce.inspect_template(flipkart_like(), "kurta.xlsx").columns}
    layout = ce.inspect_template(flipkart_like(), "kurta.xlsx",
                                 {str(col_of["Style Code"]): "parent_sku", str(col_of["Fabric"]): ""})
    by_header = {h: f for _, h, f in layout.columns}
    assert by_header["Style Code"] == "parent_sku" and by_header["Fabric"] is None


def test_a_file_that_is_not_a_template_is_refused_in_words():
    wb = Workbook(); wb.active.append(["hello", "world"])
    with pytest.raises(ValueError, match="blank upload template"):
        ce.inspect_template(_save(wb), "notes.xlsx")


def test_master_sheet_opens_with_one_header_row_and_our_rows():
    wb = load_workbook(io.BytesIO(ce.master_xlsx(_rows())))
    ws = wb["ZISUN catalogue"]
    assert ws["A1"].value == "SKU" and ws.max_row == 3
    assert ce.master_csv(_rows()).startswith("﻿".encode())
